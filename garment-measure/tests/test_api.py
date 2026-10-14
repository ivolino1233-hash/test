"""API flow: styles/specs, measure a photo, adjust a point, save, AI check, history, export."""
import io
import time

import cv2
import httpx
import numpy as np
import pytest
from fastapi.testclient import TestClient

from app import specs
from app.config import Settings
from app.main import create_app
from app.qwen import QwenClient, _json
from tests import fake_qwen
from tests.synth.cases import photo


@pytest.fixture
def client(tmp_path):
    st = Settings(data_dir=tmp_path, cert_dir=tmp_path / "certs", qwen_base_url="http://qwen.test/v1")
    app = create_app(st, qwen_transport=httpx.MockTransport(fake_qwen.handler))
    with TestClient(app) as c:
        yield c


SWEATER_SPEC = [
    {"size": "M", "key": "chest", "target_mm": 520, "tol_minus_mm": 10, "tol_plus_mm": 10},
    {"size": "M", "key": "body_length", "target_mm": 660, "tol_minus_mm": 10, "tol_plus_mm": 10},
    {"size": "M", "key": "sleeve_length", "target_mm": 585, "tol_minus_mm": 10, "tol_plus_mm": 10},
    {"size": "M", "key": "hem", "target_mm": 496, "tol_minus_mm": 10, "tol_plus_mm": 10},
]


def _jpeg(img):
    return cv2.imencode(".jpg", img, [cv2.IMWRITE_JPEG_QUALITY, 92])[1].tobytes()


def _wait_ai(client, scan_id, timeout=10):
    t = time.time()
    while time.time() - t < timeout:
        r = client.get(f"/api/scans/{scan_id}/ai").json()
        if r["ai_status"] != "pending":
            return r
        time.sleep(0.1)
    raise AssertionError("AI job did not finish")


def test_full_flow(client):
    assert client.get("/api/status").json()["qwen"]["model"] == fake_qwen.MODEL
    st = client.post("/api/styles", json={"code": "KN-1", "name": "Crew", "garment_type": "sweater",
                                          "sizes": ["M", "L"], "spec": SWEATER_SPEC}).json()
    assert len(st["spec"]) == 4
    assert client.post("/api/styles", json={"code": "kn-1", "garment_type": "sweater"}).status_code == 409

    img, _ = photo("sweater")
    r = client.post("/api/measure", files={"image": ("p.jpg", _jpeg(img), "image/jpeg")},
                    data={"style_id": str(st["id"]), "size": "M", "label": "PO-77"})
    assert r.status_code == 200, r.text
    scan = r.json()
    res = scan["results"]
    assert res["chest"]["status"] == "pass"
    assert res["sleeve_length.l"]["status"] == "fail"      # measured ~600 vs spec 585 +/- 10
    assert scan["verdict"].startswith("FAIL")
    assert scan["status"] == "draft" and scan["ai_status"] == "pending"
    for name in ("rect", "annotated", "thumb", "photo"):
        assert client.get(scan["images"][name]).headers["content-type"] == "image/jpeg"

    ai = _wait_ai(client, scan["id"])
    assert ai["ai_status"] == "done"
    assert ai["ai"]["check"]["lying_flat"] is True and "stretched" in ai["ai"]["summary"]

    # Drag the left cuff point 15 mm back toward the shoulder: sleeve now within tolerance.
    sp, ct = scan["landmarks"]["sp_l"], scan["landmarks"]["cuff_top_l"]
    s = scan["px_per_mm"]
    vec = [(sp[0] - ct[0]), (sp[1] - ct[1])]
    n = (vec[0] ** 2 + vec[1] ** 2) ** 0.5
    new_ct = [ct[0] + vec[0] / n * 15 * s, ct[1] + vec[1] / n * 15 * s]
    r = client.put(f"/api/scans/{scan['id']}", json={"landmarks": {"cuff_top_l": new_ct}, "status": "saved",
                                                     "notes": "checked"})
    up = r.json()
    assert up["results"]["sleeve_length.l"]["status"] == "pass"
    assert up["confidence"]["cuff_top_l"] == 1.0 and up["status"] == "saved"
    _wait_ai(client, scan["id"])

    hist = client.get("/api/scans").json()
    assert [h["id"] for h in hist] == [scan["id"]]
    assert client.get(f"/api/scans?style_id={st['id']}").json()[0]["label"] == "PO-77"
    csv_text = client.get("/api/export.csv").content.decode("utf-8-sig")
    assert "Chest width" in csv_text and "KN-1" in csv_text
    x = client.get("/api/export.xlsx")
    assert x.content[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Scans", "Measurements"]

    assert client.delete(f"/api/scans/{scan['id']}").json()["ok"]
    assert client.get(f"/api/scans/{scan['id']}").status_code == 404


def test_measure_error_returns_reason(client):
    img = cv2.imencode(".jpg", np.full((600, 800, 3), 180, np.uint8))[1].tobytes()
    r = client.post("/api/measure", files={"image": ("p.jpg", img, "image/jpeg")})
    assert r.status_code == 422 and r.json()["error"]["code"] == "NO_MARKER"


def test_preview_endpoint(client):
    img, _ = photo("trousers")
    small = cv2.resize(img, (1280, 960), interpolation=cv2.INTER_AREA)
    r = client.post("/api/preview", content=_jpeg(small), headers={"content-type": "image/jpeg"}).json()
    assert r["status"] == "READY" and r["outline"]


def test_settings_and_marker(client):
    s = client.put("/api/settings", json={"ruler_h_mm": 148.5, "unit": "in"}).json()
    assert s["ruler_h_mm"] == 148.5 and s["unit"] == "in"
    assert client.put("/api/settings", json={"ruler_h_mm": 100}).status_code == 400
    pdf = client.get("/marker.pdf?sheets=2")
    assert pdf.content.startswith(b"%PDF")
    assert client.get("/api/qr.svg?text=https://192.168.1.5:8443").content.startswith(b"<?xml")
    assert client.get("/").status_code == 200 and client.get("/onboarding").status_code == 200


def test_spec_csv_import(client):
    up = client.post("/api/styles/import", files={"file": ("s.csv", specs.template_csv().encode(), "text/csv")})
    body = up.json()
    assert body["imported"] == ["KN-1001"] and not body["errors"]
    st = client.get("/api/styles").json()[0]
    full = client.get(f"/api/styles/{st['id']}").json()
    assert full["sizes"] == ["M", "L"]
    keys = {(p["size"], p["key"]) for p in full["spec"]}
    assert ("M", "chest") in keys and ("L", "sleeve_opening") in keys
    chest_m = next(p for p in full["spec"] if p["size"] == "M" and p["key"] == "chest")
    assert chest_m["target_mm"] == 520 and chest_m["tol_minus_mm"] == 10


def test_bad_csv_rows_are_reported():
    text = "style_code,size,point,target_cm\nA,M,elbow width,10\nA,M,1/2 chest,abc\n"
    parsed, errors = specs.parse_spec_csv(text)
    assert len(errors) == 2 and "elbow width" in errors[0]


def test_pin_required(tmp_path):
    st = Settings(data_dir=tmp_path, cert_dir=tmp_path, qwen_enabled=False, app_pin="4321")
    with TestClient(create_app(st)) as c:
        assert c.get("/api/styles").status_code == 401
        assert c.get("/api/status").json()["pin_required"] is True
        assert c.post("/api/login", json={"pin": "0000"}).status_code == 401
        assert c.post("/api/login", json={"pin": "4321"}).status_code == 200
        assert c.get("/api/styles").status_code == 200


def test_qwen_unreachable_does_not_block(tmp_path):
    def down(request):
        raise httpx.ConnectError("connection refused")

    st = Settings(data_dir=tmp_path, cert_dir=tmp_path, qwen_base_url="http://qwen.test/v1")
    with TestClient(create_app(st, qwen_transport=httpx.MockTransport(down))) as c:
        assert c.get("/api/status").json()["qwen"]["reachable"] is False
        img, _ = photo("skirt")
        r = c.post("/api/measure", files={"image": ("p.jpg", _jpeg(img), "image/jpeg")})
        assert r.status_code == 200
        ai = _wait_ai(c, r.json()["id"])
        assert ai["ai_status"] == "error"


def test_qwen_json_parsing_variants():
    assert _json('<think>hmm</think>```json\n{"a": 1}\n```') == {"a": 1}
    assert _json('Sure! {"a": 2} hope this helps') == {"a": 2}


@pytest.mark.anyio
async def test_qwen_falls_back_without_structured_output():
    def h(request):
        import json
        body = json.loads(request.content) if request.content else {}
        if request.url.path.endswith("/models"):
            return httpx.Response(200, json={"data": [{"id": "qwen-local"}]})
        if "response_format" in body:
            return httpx.Response(400, json={"error": "response_format not supported"})
        return httpx.Response(200, json=fake_qwen.reply(body))

    q = QwenClient("http://x/v1", transport=httpx.MockTransport(h))
    out = await q.photo_check(b"\xff\xd8fake", "Sweater")
    assert out["lying_flat"] is True and q.supports_schema is False
    await q.close()
