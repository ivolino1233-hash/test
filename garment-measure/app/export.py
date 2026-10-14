"""History export: one row per scan per measurement point (CSV, or Excel with a summary sheet)."""
from __future__ import annotations

import csv
import io

from .specs import verdict

COLUMNS = ["scan_id", "date", "style_code", "style_name", "size", "label", "garment_type", "point", "side",
           "spec_cm", "tol_minus_cm", "tol_plus_cm", "actual_cm", "diff_cm", "status"]


def _cm(v):
    return None if v is None else round(v / 10, 2)


def rows(scans: list[dict], styles: dict) -> list[list]:
    out = []
    for s in scans:
        st = styles.get(s["style_id"]) or {}
        for r in (s["results"] or {}).values():
            out.append([s["id"], s["created_at"][:19].replace("T", " "), st.get("code", ""), st.get("name", ""),
                        s["size"], s["label"], s["garment_type"], r["label"], {"l": "left", "r": "right"}.get(
                            r.get("side") or "", ""), _cm(r["target_mm"]), _cm(r["tol_minus_mm"]),
                        _cm(r["tol_plus_mm"]), _cm(r["actual_mm"]), _cm(r["diff_mm"]), r["status"]])
    return out


def to_csv(scans, styles) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS)
    w.writerows(rows(scans, styles))
    return ("﻿" + buf.getvalue()).encode("utf-8")   # BOM so Excel opens it as UTF-8


def to_xlsx(scans, styles) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Scans"
    ws.append(["scan_id", "date", "style_code", "size", "label", "garment_type", "passed", "failed", "verdict"])
    for s in scans:
        st = styles.get(s["style_id"]) or {}
        ws.append([s["id"], s["created_at"][:19].replace("T", " "), st.get("code", ""), s["size"], s["label"],
                   s["garment_type"], s["pass_count"], s["fail_count"], verdict(s["results"] or {})])
    ws2 = wb.create_sheet("Measurements")
    ws2.append(COLUMNS)
    red, green = PatternFill("solid", fgColor="F8D7DA"), PatternFill("solid", fgColor="D4EDDA")
    for row in rows(scans, styles):
        ws2.append(row)
        if row[-1] in ("pass", "fail"):
            ws2.cell(ws2.max_row, len(COLUMNS)).fill = green if row[-1] == "pass" else red
    for sheet in (ws, ws2):
        for c in sheet[1]:
            c.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col[:200])
            sheet.column_dimensions[col[0].column_letter].width = min(48, max(9, width + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
